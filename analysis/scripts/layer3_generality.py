#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Layer 3 - Generality benchmark for the projection artefact.

WHY THIS LAYER EXISTS
---------------------
Stages A and B both tested ONE signature (the stiffness-Hippo programme derived
in Layer 0). Even a perfectly concordant pair of nulls cannot distinguish

    (a) "this particular signature does not transfer"
from
    (b) "the projection METHOD itself is not tissue-specific".

To make claim (b) - which is the claim the manuscript rests on - five
INDEPENDENTLY PUBLISHED ECM / mechanotransduction gene sets are pushed through
the identical scoring and null machinery. The method-level claim is sustained
only if most of them fail specificity IN THE SAME WAY.

WHAT IS ACTUALLY TESTED
-----------------------
Not activation. SPECIFICITY:

    does set g shift further in the target pain tissue (DRG, SNI)
    than it does in a tissue that cannot plausibly host a stiffness
    programme (blood, SNI)?

Specificity index SI_g = |MRU(DRG,SNI) - 0.5| / |MRU(BLD,SNI) - 0.5|.
SI_g < 1  ->  blood shifts at least as far as DRG  ->  specificity FAILED.

FROZEN BEFORE EXECUTION
-----------------------
Gene-set membership, the statistic (MRU), the null construction
(expressed-only background, baseline decile stratification, WITHOUT
replacement), the iteration count, the seed, the multiplicity family
(5 sets x 4 tissues x 2 models = 40 tests) and the decision rule
(>= 4 of 5 fail) were all written into protocol v0.5 / ms_panel.json
before any Layer-3 output was inspected.

Outputs (analysis/results/layer3/):
  layer3_generality.json   per-set endpoint table, SI, and the primary verdict
  layer3_generality.csv    flat table, one row per (set x tissue x model)
"""

import json
import os
import sys

import numpy as np
import pandas as pd

# ---------------------------------------------------------------- paths ----
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
CFG_PATH = os.path.join(ROOT, "analysis", "config", "ms_panel.json")
RAW_DIR = os.path.join(ROOT, "analysis", "data", "raw")
OUT_DIR = os.path.join(ROOT, "analysis", "results", "layer3")
os.makedirs(OUT_DIR, exist_ok=True)

TISSUES = ["DRG", "SC", "BRN", "BLD"]
MODELS = ["CFA", "SNI"]
TABLES = {
    ("BRN", "CFA"): "GSE111216_Parisien_Suppl_Table_01_gene_BRN_CTRvsCFA.xlsx",
    ("BRN", "SNI"): "GSE111216_Parisien_Suppl_Table_02_gene_BRN_CTRvsSNI.xlsx",
    ("SC", "CFA"): "GSE111216_Parisien_Suppl_Table_03_gene_SC_CTRvsCFA.xlsx",
    ("SC", "SNI"): "GSE111216_Parisien_Suppl_Table_04_gene_SC_CTRvsSNI.xlsx",
    ("DRG", "CFA"): "GSE111216_Parisien_Suppl_Table_05_gene_DRG_CTRvsCFA.xlsx",
    ("DRG", "SNI"): "GSE111216_Parisien_Suppl_Table_06_gene_DRG_CTRvsSNI.xlsx",
    ("BLD", "CFA"): "GSE111216_Parisien_Suppl_Table_07_gene_BLD_CTRvsCFA.xlsx",
    ("BLD", "SNI"): "GSE111216_Parisien_Suppl_Table_08_gene_BLD_CTRvsSNI.xlsx",
}


def human_to_mouse(sym):
    """Frozen capitalisation rule (identical to layer1_projection.human_to_mouse)."""
    if sym.isupper():
        return sym[0] + sym[1:].lower()
    return sym


def bh(p):
    """Benjamini-Hochberg q-values."""
    p = np.asarray(p, dtype=float)
    n = len(p)
    o = np.argsort(p)
    q = p[o] * n / (np.arange(n) + 1)
    q = np.minimum.accumulate(q[::-1])[::-1]
    out = np.empty(n)
    out[o] = q
    return np.clip(out, 0, 1)


def ranks_of(df, var):
    """0-based ranks ascending by `var` (ties broken by mergesort stability)."""
    order = np.argsort(df[var].to_numpy(), kind="mergesort")
    arr = np.empty(len(df))
    arr[order] = np.arange(len(df))
    return pd.Series(arr, index=df.index)


def load_table(path):
    """Read one deposited per-contrast xlsx into a DataFrame indexed by gene symbol."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    data = [r for r in rows if r is not None and not all(v is None for v in r)]
    df = pd.DataFrame(data, columns=header)
    wb.close()
    df = df.rename(columns={"gene": "symbol"})
    for c in ["value_CTR", "p_value", "log2_fc", "log2_fcZ", "pi_value"]:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    df["symbol"] = df["symbol"].astype(str).str.strip()
    df = df[~df["symbol"].isin(["", "nan", "None"])]
    # BUG FIX: de-duplicate on SYMBOL before it becomes the index. Filtering on
    # `df.index` at this point targets the default RangeIndex, which is always
    # unique, so duplicated symbols survived and made `.loc[pick]` return more
    # rows than requested (broadcast error) - and would silently misalign the
    # permutation null if the broadcast happened to line up.
    df = df.drop_duplicates(subset="symbol", keep="first")
    return df.set_index("symbol")


def _stratified_draws(rk, member_bins, pool_by_bin, n_perm, rng, block=1000):
    """Draw a DISTINCT gene per member, each from that member's baseline decile.

    rk is indexed by symbol, so pool entries are SYMBOLS and must be looked up
    with .loc (indexing the values array with strings raises IndexError).

    Distinctness is enforced by construction rather than by rejection sampling.
    Decile pools are disjoint, so two members can only collide if they share a
    decile; within a decile we draw k items WITHOUT replacement using random
    keys plus argpartition, which is O(pool) per row. The original rejection
    sampler did an O(m^2) pairwise scan per retry and made a 5 x 8 x 10,000
    design intractable.
    """
    m = len(member_bins)
    out = np.empty((n_perm, m), dtype=float)

    # member positions grouped by decile
    by_bin = {}
    for j, b in enumerate(member_bins):
        by_bin.setdefault(b, []).append(j)

    # rank lookups are cheapest when done on the numpy values array, so cache
    # a symbol -> row position index once.
    pos = pd.Series(np.arange(len(rk)), index=rk.index)

    for b, positions in by_bin.items():
        pool = pool_by_bin[b]
        need = len(positions)
        if need > len(pool):
            # degenerate decile: cannot draw enough distinct genes
            pick = pool[rng.integers(0, len(pool), size=(n_perm, need))]
            pick_pos = pos.loc[pick.ravel()].to_numpy().reshape(pick.shape)
        else:
            pick_pos = np.empty((n_perm, need), dtype=np.int64)
            start = 0
            while start < n_perm:
                stop = min(start + block, n_perm)
                keys = rng.random((stop - start, len(pool)))
                idx = np.argpartition(keys, need - 1, axis=1)[:, :need]
                pick_pos[start:stop] = pos.loc[pool[idx.ravel()]].to_numpy(
                    dtype=np.int64).reshape(idx.shape)
                start = stop
        vals = rk.to_numpy()
        for a, j in enumerate(positions):
            out[:, j] = vals[pick_pos[:, a]]

    return out


def _without_replacement(n_perm, m, n, rng, block=2000, max_retry=25):
    """n_perm rows of m DISTINCT indices from [0, n)."""
    idx = rng.integers(0, n, size=(n_perm, m))
    for _ in range(max_retry):
        dup = np.array([np.unique(r).size < m for r in idx])
        if not dup.any():
            break
        idx[dup] = rng.integers(0, n, size=(int(dup.sum()), m))
    return idx


def main():
    cfg = json.load(open(CFG_PATH, encoding="utf-8"))
    l3 = cfg["layer3"]
    rt = cfg["layer1"]["rank_test"]
    rank_var = rt["primary_ranking_variable"]
    n_perm = l3["null"]["n_iterations"]
    seed = l3["null"]["seed"]
    rng = np.random.default_rng(seed)

    tgt = l3["target_contrast"]
    neg = l3["negative_contrast"]

    rows = []
    per_set = {}

    for gname, gdef in l3["gene_sets"].items():
        mouse_members = [human_to_mouse(g) for g in gdef["members"]]
        per_set[gname] = {"source": gdef["source"],
                          "n_human": len(gdef["members"]),
                          "mouse_members": mouse_members}

        for tis in TISSUES:
            for mod in MODELS:
                path = os.path.join(RAW_DIR, TABLES[(tis, mod)])
                df = load_table(path)

                # expressed-only background: drop baseline-zero loci
                n_all = len(df)
                expr = df[df["value_CTR"] > 0]
                n_expr = len(expr)
                expr = expr[np.isfinite(expr[rank_var])]
                rk = ranks_of(expr, rank_var)
                n_rank = len(expr)

                present = [m for m in mouse_members if m in expr.index]
                if len(present) < 3:
                    print(f"[L3] {gname} {tis}-{mod}: only {len(present)} "
                          f"members present - skipped")
                    continue

                obs = float(np.median(rk.loc[present].to_numpy()) / (n_rank - 1))

                # stratified (size + baseline decile) null, without replacement
                bins = pd.qcut(np.log2(expr["value_CTR"] + 1),
                               q=10, labels=False, duplicates="drop")
                pool_by_bin = {}
                for b in np.unique(bins):
                    pool_by_bin[b] = expr.index[bins == b].to_numpy()
                member_bins = [int(bins.loc[m]) for m in present]

                perm = _stratified_draws(rk, member_bins, pool_by_bin,
                                         n_perm, rng)
                perm_mru = np.median(perm, axis=1) / (n_rank - 1)
                p_matched = (int((np.abs(perm_mru - 0.5)
                                  >= abs(obs - 0.5)).sum()) + 1) / (n_perm + 1)

                # unstratified expressed-only null, for contrast
                idx = _without_replacement(n_perm, len(present), n_rank, rng)
                perm2 = np.median(rk.to_numpy()[idx], axis=1) / (n_rank - 1)
                p_unmatched = (int((np.abs(perm2 - 0.5)
                                    >= abs(obs - 0.5)).sum()) + 1) / (n_perm + 1)

                rows.append({
                    "gene_set": gname, "tissue": tis, "model": mod,
                    "n_human": len(gdef["members"]),
                    "n_mouse_present": len(present),
                    "n_expressed_background": n_rank,
                    "n_rows_total": n_all,
                    "MRU": obs,
                    "abs_shift": abs(obs - 0.5),
                    "p_matched": p_matched,
                    "p_unmatched": p_unmatched,
                })

    tbl = pd.DataFrame(rows)
    tbl["q_matched_BH40"] = bh(tbl["p_matched"].to_numpy())

    # ---- specificity index ------------------------------------------------
    def mru(g, tis, mod):
        s = tbl[(tbl.gene_set == g) & (tbl.tissue == tis) & (tbl.model == mod)]
        return float(s["MRU"].iloc[0]) if len(s) else float("nan")

    spec = {}
    for g in l3["gene_sets"]:
        st, sn = mru(g, tgt["tissue"], tgt["model"]), mru(g, neg["tissue"], neg["model"])
        si = (abs(st - 0.5) / abs(sn - 0.5)) if (sn is not None and abs(sn - 0.5) > 0) else float("nan")
        spec[g] = {
            "MRU_target": st,
            "MRU_negative_blood": sn,
            "SI": si,
            "specificity_failed": bool(si < 1) if si == si else None,
        }

    n_fail = sum(1 for v in spec.values() if v["specificity_failed"])
    n_sets = len(spec)
    primary_pass = bool(n_fail >= 4)

    print("=" * 78)
    print("LAYER 3 - generality of the projection artefact")
    print("=" * 78)
    print(f"null: {n_perm} perms, seed {seed}, expressed-only background, "
          f"baseline-decile stratified, WITHOUT replacement\n")
    for g, v in spec.items():
        print(f"  {g:34s} MRU(DRG,SNI)={v['MRU_target']:.3f}  "
              f"MRU(BLD,SNI)={v['MRU_negative_blood']:.3f}  "
              f"SI={v['SI']:.3f}  fail={v['specificity_failed']}")
    print(f"\n  sets failing specificity: {n_fail}/{n_sets}")
    print(f"  PRIMARY (L3-generality, need >= 4/5): "
          f"{'PASS' if primary_pass else 'FAIL'}")
    print(f"\n  BH q<=0.05 across 40 tests: "
          f"{int((tbl['q_matched_BH40'] <= 0.05).sum())}")

    out = {
        "protocol": cfg["_protocol"],
        "design": {
            "statistic": l3["statistic"],
            "null": l3["null"],
            "specificity_index": l3["specificity_index"],
            "primary_rule": l3["primary_endpoint"]["rule"],
            "target_contrast": tgt,
            "negative_contrast": neg,
        },
        "per_set_specificity": spec,
        "n_sets_failing_specificity": n_fail,
        "n_sets_tested": n_sets,
        "primary_endpoint_pass": primary_pass,
        "n_tests": int(len(tbl)),
        "n_surviving_BH_0.05": int((tbl["q_matched_BH40"] <= 0.05).sum()),
        "min_q_matched": float(tbl["q_matched_BH40"].min()),
        "table": tbl.to_dict(orient="records"),
    }
    with open(os.path.join(OUT_DIR, "layer3_generality.json"), "w",
              encoding="utf-8") as fh:
        json.dump(out, fh, indent=2, ensure_ascii=False)
    tbl.to_csv(os.path.join(OUT_DIR, "layer3_generality.csv"), index=False)
    print(f"\nwrote {OUT_DIR}/layer3_generality.json")
    print(f"wrote {OUT_DIR}/layer3_generality.csv")


if __name__ == "__main__":
    main()
