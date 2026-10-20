#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AUDIT ROUND 1b - separate MAPPING failure from COVERAGE loss.

Protocol section 4.3 defines mapping_rate = matched / |signature| and attaches a
governance rule (rate < 0.7 -> primary endpoint downgraded to exploratory).

Stage B, however, implicitly folded TWO different quantities into that one number:
  1. cross-species homology mapping   (human symbol -> mouse symbol present in the
                                       annotation carried over from GSE111216)
  2. coverage after the CPM>1 in >=4/16 expression filter of the DRG matrix

Only (1) is a mapping failure. (2) is a power / biology matter: a signature member
that is simply not expressed in DRG. Conflating them can wrongly invalidate a
designated primary endpoint.

This script reports both, separately, per signature, and states the governance
consequence of each.
"""

import gzip
import json
import os
import re
import sys

import numpy as np
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
CFG = os.path.join(ROOT, "analysis", "config", "ms_panel.json")
RAW = os.path.join(ROOT, "analysis", "data", "raw")
L0 = os.path.join(ROOT, "analysis", "results", "layer0")
OUT = os.path.join(ROOT, "analysis", "results", "audit")
os.makedirs(OUT, exist_ok=True)
sys.path.insert(0, HERE)

FNAME_RE = re.compile(r"^(GSM\d+)_(DMI|saline)-(sham|SNI)(\d)_S\d+\.counts\.txt\.gz$",
                      re.IGNORECASE)


def human_to_mouse(sym):
    if sym.isupper():
        return sym[0] + sym[1:].lower()
    return sym


def load_symbol_map():
    wb = openpyxl.load_workbook(
        os.path.join(RAW, "GSE111216_Parisien_Suppl_Table_06_gene_DRG_CTRvsSNI.xlsx"),
        read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows)]
    gi, gs = header.index("gene_id"), header.index("gene")
    m = {}
    for r in rows:
        if r is None or r[gi] is None or r[gs] is None:
            continue
        m.setdefault(str(r[gi]).split(".")[0], str(r[gs]).strip())
    wb.close()
    return m


def build_counts():
    cnt = os.path.join(RAW, "_gse261676")
    files = sorted(f for f in os.listdir(cnt) if FNAME_RE.match(f))
    meta, series = {}, {}
    for fn in files:
        m = FNAME_RE.match(fn)
        gsm, agent_raw, group, rep = m.group(1), m.group(2), m.group(3), m.group(4)
        agent = "DMI" if agent_raw.lower() == "dmi" else "saline"
        grp = "SNI" if group.lower() == "sni" else "sham"
        label = f"{agent}_{grp}{rep}"
        counts = {}
        with gzip.open(os.path.join(cnt, fn), "rt") as fh:
            for line in fh:
                if not line.strip():
                    continue
                parts = line.rstrip("\n").split("\t")
                if len(parts) < 2:
                    continue
                counts[parts[0].split(".")[0]] = float(parts[1])
        series[label] = pd.Series(counts)
        meta[label] = {"gsm": gsm, "agent": agent, "group": grp, "rep": int(rep)}
    return pd.DataFrame(series), pd.DataFrame(meta).T


def main():
    cfg = json.load(open(CFG, encoding="utf-8"))
    sig = json.load(open(os.path.join(L0, "layer0_signature.json"), encoding="utf-8"))
    min_rate = cfg["layer1"]["species_mapping"]["min_acceptable_mapping_rate"]

    cmat, meta = build_counts()
    lib = cmat.sum(axis=0)
    cpm = cmat / lib * 1e6
    keep = (cpm > 1).sum(axis=1) >= 4

    smap = load_symbol_map()

    # universe of mouse symbols BEFORE the expression filter
    all_syms = set(smap.get(g) for g in cmat.index)
    all_syms.discard(None)
    # universe AFTER the expression filter
    kept_syms = set(smap.get(g) for g in cpm.index[keep])
    kept_syms.discard(None)

    rep = {}
    print(f"[AUD1b] mouse symbols in annotation          : {len(all_syms)}")
    print(f"[AUD1b] mouse symbols after CPM filter       : {len(kept_syms)}")
    print(f"[AUD1b] protocol mapping floor               : {min_rate}\n")

    rows = []
    for key in ["hippo_dependent_up", "hippo_dependent_down",
                "hippo_independent_up", "hippo_independent_down"]:
        members = sig["sets"].get(key, [])
        if not members:
            continue
        mapped = [human_to_mouse(g) for g in members]
        in_annot = [m for m in mapped if m in all_syms]
        in_filt = [m for m in mapped if m in kept_syms]
        map_rate = len(in_annot) / len(members)
        cov_rate = len(in_filt) / len(members)
        lost_annot = [m for m in mapped if m not in all_syms]
        lost_filt = [m for m in in_annot if m not in kept_syms]
        rows.append({
            "signature": key,
            "n_human": len(members),
            "n_mapped_in_annotation": len(in_annot),
            "mapping_rate": round(map_rate, 4),
            "n_after_expression_filter": len(in_filt),
            "coverage_rate": round(cov_rate, 4),
            "lost_to_missing_annotation": sorted(lost_annot),
            "lost_to_expression_filter": sorted(lost_filt),
            "governance_verdict": (
                "MAPPING OK - coverage loss is a power issue, "
                "primary-endpoint status preserved"
                if map_rate >= min_rate else
                "MAPPING BREACH - downgrade to exploratory"
            ),
        })
        print(f"  {key}")
        print(f"     mapping_rate (homology)      = {len(in_annot)}/{len(members)} = {map_rate:.3f}")
        print(f"     coverage_rate (expressed)    = {len(in_filt)}/{len(members)} = {cov_rate:.3f}")
        if lost_annot:
            print(f"     lost: annotation             = {lost_annot}")
        if lost_filt:
            print(f"     lost: not expressed in DRG   = {lost_filt}")
        print(f"     verdict: {rows[-1]['governance_verdict']}\n")

    rep["per_signature"] = rows
    rep["_note"] = (
        "The protocol's single mapping_rate number conflated cross-species homology "
        "mapping with post-filter coverage. They are reported separately here. The "
        "governance rule (section 4.3) applies to the homology mapping rate only; "
        "coverage loss is a power limitation to be declared in Limitations."
    )

    with open(os.path.join(OUT, "audit_round1b.json"), "w", encoding="utf-8") as fh:
        json.dump({"_tag": cfg["_tag"], "_protocol": cfg["_protocol"], "audit": rep},
                  fh, indent=2, ensure_ascii=False)
    print("[AUD1b] wrote audit_round1b.json")


if __name__ == "__main__":
    main()
